"""Tests for build_workbook.py.

The load-bearing one is test_signals_are_not_capped: flattening to signal1/2/3 silently
loses the fourth, which is why the workbook keeps them long.
"""
import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
SCRIPT = REPO / "skills" / "find-contacts" / "scripts" / "build_workbook.py"
spec = importlib.util.spec_from_file_location("build_workbook", SCRIPT)
bw = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bw)


def a_record(**over):
    rec = {
        "company": "Ostvale Provisions", "domain": "ostvale.example.com",
        "country": "Netherlands", "staff": 180, "company_fit": 5,
        "campaign_relevance": 5, "score": 5, "routing": "white-glove",
        "liveness": "verified-200", "status": "ok",
        "people": [{"recipient_id": "R-0001", "full_name": "Priya Raman",
                    "title": "VP Supply Chain", "primary": "yes",
                    "named_by_customer": "no", "confidence": "high"}],
        "signals": [{"recipient_id": "R-0001", "fact": f"Event {i}", "date": f"2026-0{i}",
                     "angle": f"Angle {i}"} for i in range(1, 5)],
        "open_items": [{"item": "Dates disagree", "affects": "Ostvale",
                        "status": "check before use", "detail": "Newsroom undated."}],
    }
    rec.update(over)
    return rec


def build(tmp_path, records):
    out = tmp_path / "wb.xlsx"
    counts = bw.build(records, out)
    return load_workbook(out), counts


def test_four_sheets_in_order(tmp_path):
    wb, _ = build(tmp_path, [a_record()])
    assert wb.sheetnames == ["Companies", "Contacts", "Signals", "Open items"]


def test_signals_are_not_capped(tmp_path):
    """Four signals in, four rows out. Flattening would lose the fourth."""
    wb, counts = build(tmp_path, [a_record()])
    assert counts["Signals"] == 4
    assert wb["Signals"].max_row - 1 == 4


def test_scoring_and_routing_reach_the_sheet(tmp_path):
    wb, _ = build(tmp_path, [a_record()])
    hdr = [c.value for c in wb["Companies"][1]]
    for col in ("company_fit", "campaign_relevance", "score", "routing", "staff"):
        assert col in hdr, f"{col} missing from Companies"
    row = {h: c.value for h, c in zip(hdr, wb["Companies"][2], strict=False)}
    assert row["routing"] == "white-glove"
    assert row["score"] == 5


def test_contact_judgement_fields_survive(tmp_path):
    wb, _ = build(tmp_path, [a_record()])
    hdr = [c.value for c in wb["Contacts"][1]]
    for col in ("primary", "named_by_customer", "flags", "role_status"):
        assert col in hdr, f"{col} missing from Contacts"


def test_open_items_get_their_own_sheet(tmp_path):
    wb, counts = build(tmp_path, [a_record()])
    assert counts["Open items"] == 1
    assert wb["Open items"]["A2"].value == "Dates disagree"


def test_a_company_with_no_people_still_appears(tmp_path):
    """A dropped or thin company must not vanish from the list."""
    wb, counts = build(tmp_path, [a_record(people=[], signals=[], routing="drop")])
    assert counts["Companies"] == 1
    assert counts["Contacts"] == 0


def test_flat_csv_is_lossy_and_says_so(tmp_path):
    out = tmp_path / "flat.csv"
    bw.to_csv([a_record()], out)
    hdr = out.read_text(encoding="utf-8").splitlines()[0]
    assert "signal3_angle" in hdr and "signal4" not in hdr
    assert "perso_1" in hdr and "perso_4" in hdr


def test_control_characters_are_stripped(tmp_path):
    rec = a_record()
    rec["fit_reason"] = "Own\x07 fleet\x00 and planners."
    wb, _ = build(tmp_path, [rec])
    hdr = [c.value for c in wb["Companies"][1]]
    row = {h: c.value for h, c in zip(hdr, wb["Companies"][2], strict=False)}
    assert row["fit_reason"] == "Own fleet and planners."


def test_cli_reports_counts(tmp_path):
    rp = tmp_path / "r.json"
    rp.write_text(json.dumps([a_record()]), encoding="utf-8")
    r = subprocess.run([sys.executable, str(SCRIPT), str(rp), str(tmp_path / "o.xlsx")],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert "4 signals" in r.stdout


def test_missing_file_exits_2(tmp_path):
    r = subprocess.run([sys.executable, str(SCRIPT), str(tmp_path / "no.json"),
                        str(tmp_path / "o.xlsx")], capture_output=True, text=True)
    assert r.returncode == 2


# --------------------------------------------------------------- field loss

def test_unwritten_fields_are_reported_not_swallowed():
    """The defect this guard exists for. Three live runs lost a field this way - a
    published email, a linkedin_url a person-anchored provider needed, and a resolved
    employer flag - and every time the output looked complete."""
    parts = bw.unpack([{
        "company": "Ostvale", "domain": "ostvale.example.com",
        "people": [{"recipient_id": "R-1", "full_name": "A One",
                    "email": "a@ostvale.example.com", "invented_field": "x"}],
        "signals": [{"fact": "f", "another_invention": "y"}],
        "open_items": [],
    }])
    missing = bw.unwritten_fields(parts)
    assert missing["Contacts"] == ["invented_field"]
    assert missing["Signals"] == ["another_invention"]


def test_structural_keys_are_not_reported_as_lost():
    """people/signals/open_items become their own sheets. Reporting them as dropped
    would be noise that trains the reader to ignore the warning."""
    parts = bw.unpack([{"company": "Ostvale", "people": [], "signals": [],
                        "open_items": []}])
    assert bw.unwritten_fields(parts) == {}


def test_declared_contract_fields_are_never_reported():
    parts = bw.unpack([{
        "company": "Ostvale", "domain": "d", "country": "NL", "industry": "Food",
        "audience": "distributor", "segment": "distributor", "score": 5,
        "people": [{"recipient_id": "R-1", "full_name": "A One", "title": "VP",
                    "email": "a@b.c", "linkedin_url": "https://x", "flags": ""}],
        "signals": [{"recipient_id": "R-1", "fact": "f", "date": "2026-05",
                     "date_confidence": "month", "type": "expansion",
                     "for_person": "ALL", "scope": "company-wide",
                     "source_url": "https://x", "angle": "a"}],
        "open_items": [{"item": "i", "affects": "x", "status": "open", "detail": "d"}],
    }])
    assert bw.unwritten_fields(parts) == {}


def test_signals_sheet_carries_for_person_and_scope():
    """Contract fields that never reach a column are the same bug in slower motion."""
    declared = {n for n, _ in bw.SHEETS["Signals"]}
    assert {"for_person", "scope"} <= declared


def test_companies_sheet_carries_audience():
    """audience is what the outreach profile matches sequence variants against. A
    reviewer who cannot see it cannot tell why a given variant was chosen."""
    assert "audience" in {n for n, _ in bw.SHEETS["Companies"]}


def test_the_shipped_example_loses_nothing(tmp_path):
    """Whatever ships as the worked example has to survive its own builder."""
    data = json.loads((REPO / "examples" / "enriched.example.json").read_text())
    assert bw.unwritten_fields(bw.unpack(data)) == {}


def test_build_fails_loudly_on_an_unwritten_field(tmp_path, capsys):
    data = [{"company": "Ostvale", "people": [{"recipient_id": "R-1", "surprise": "z"}],
             "signals": [], "open_items": []}]
    with pytest.raises(SystemExit) as exc:
        bw.build(data, tmp_path / "out.xlsx", strict_fields=True)
    assert exc.value.code == 2
    assert "surprise" in capsys.readouterr().err


def test_build_warns_but_proceeds_when_loss_is_allowed(tmp_path, capsys):
    data = [{"company": "Ostvale", "people": [{"recipient_id": "R-1", "surprise": "z"}],
             "signals": [], "open_items": []}]
    counts = bw.build(data, tmp_path / "out.xlsx", strict_fields=False)
    assert counts["Contacts"] == 1
    assert "surprise" in capsys.readouterr().err
